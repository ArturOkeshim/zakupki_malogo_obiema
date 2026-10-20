"""
Выгрузка в Excel: тендеры, по которым есть оценка цен (Claude или Yandex), и все их позиции.
В сводке — fit_status: profitable / unprofitable / unknown.

Запуск:
    python export_profitable_excel.py
    python export_profitable_excel.py --source yandex
"""

import argparse
import json
from datetime import datetime, timezone
from pathlib import Path

from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from db.models import AuctionRecord, PurchaseRecord, SpecificationItemRecord
from db.session import SessionLocal
from parsing import parse_deadline_from_date_range
from price_evaluate import line_cost_for_item, tender_line_total

PROJECT_ROOT = Path(__file__).resolve().parent
DEFAULT_OUTPUT_CLAUDE = PROJECT_ROOT / "data" / "tenders_claude.xlsx"
DEFAULT_OUTPUT_YANDEX = PROJECT_ROOT / "data" / "tenders_yandex.xlsx"


def excel_agent_log(text: str | None) -> str | None:
    """Excel не принимает ячейку длиннее 32767 символов."""
    if text is None:
        return None
    limit = 32000
    if len(text) <= limit:
        return text
    return text[:limit] + "\n… обрезано для Excel"


def fit_status_label(fit_status: str | None) -> str:
    if fit_status == "profitable":
        return "Подходит"
    if fit_status == "unprofitable":
        return "Не подходит"
    return "Неполные данные"


def item_has_price_data(item: SpecificationItemRecord, source: str) -> bool:
    if source == "yandex":
        return item.yandex_research_status in ("done", "not_found")
    return item.price_research_status in ("priced", "not_found")


def auction_has_price_data(items: list[SpecificationItemRecord], source: str) -> bool:
    for item in items:
        if item_has_price_data(item, source):
            return True
    return False


def float_or_none(value) -> float | None:
    if value is None:
        return None
    return float(value)


def sources_to_urls(market_sources_json: str | None) -> str:
    if not market_sources_json:
        return ""

    try:
        sources = json.loads(market_sources_json)
    except json.JSONDecodeError:
        return ""

    urls = []
    for source in sources:
        if not isinstance(source, dict):
            continue
        url = source.get("url")
        if url:
            urls.append(str(url))
    return "; ".join(urls)


def get_latest_purchases_by_auction(session) -> dict[str, PurchaseRecord]:
    purchases = session.scalars(select(PurchaseRecord)).all()
    latest: dict[str, PurchaseRecord] = {}
    for purchase in purchases:
        current = latest.get(purchase.auction_number)
        if current is None or purchase.scraped_at > current.scraped_at:
            latest[purchase.auction_number] = purchase
    return latest


def load_tenders_for_export(
    source: str,
    auction_numbers: list[str] | None = None,
) -> tuple[list[AuctionRecord], dict[str, list[SpecificationItemRecord]], dict[str, PurchaseRecord]]:
    """
    Тендеры с detail_status=done, у которых хотя бы одна позиция уже оценена (не pending).
    """
    with SessionLocal() as session:
        all_done = session.scalars(
            select(AuctionRecord).where(AuctionRecord.detail_status == "done")
        ).all()
        if auction_numbers is not None:
            allowed = set(auction_numbers)
            all_done = [auction for auction in all_done if auction.auction_number in allowed]

        if not all_done:
            return [], {}, {}

        all_numbers = [a.auction_number for a in all_done]
        all_items = session.scalars(
            select(SpecificationItemRecord)
            .where(SpecificationItemRecord.auction_number.in_(all_numbers))
            .order_by(
                SpecificationItemRecord.auction_number,
                SpecificationItemRecord.position_index,
            )
        ).all()

        items_by_auction: dict[str, list[SpecificationItemRecord]] = {num: [] for num in all_numbers}
        for item in all_items:
            items_by_auction[item.auction_number].append(item)

        auctions_with_prices: list[AuctionRecord] = []
        for auction in all_done:
            items = items_by_auction.get(auction.auction_number, [])
            if auction_has_price_data(items, source):
                auctions_with_prices.append(auction)

        def sort_margin(auction: AuctionRecord) -> float:
            if source == "yandex":
                value = auction.margin_rub_yandex
            else:
                value = auction.margin_rub
            if value is None:
                return -1e18
            return float(value)

        auctions_with_prices.sort(key=sort_margin, reverse=True)

        export_numbers = {a.auction_number for a in auctions_with_prices}
        items_by_auction = {
            num: items for num, items in items_by_auction.items() if num in export_numbers
        }

        latest_purchases = get_latest_purchases_by_auction(session)

        for auction in auctions_with_prices:
            session.expunge(auction)
        for items in items_by_auction.values():
            for item in items:
                session.expunge(item)

        return list(auctions_with_prices), items_by_auction, latest_purchases


def autosize_columns(sheet, max_width: int = 60) -> None:
    for column_index, column_cells in enumerate(sheet.columns, start=1):
        max_length = 0
        for cell in column_cells:
            if cell.value is None:
                continue
            text = str(cell.value)
            if len(text) > max_length:
                max_length = len(text)
        width = min(max(max_length + 2, 10), max_width)
        sheet.column_dimensions[get_column_letter(column_index)].width = width


def write_summary_sheet(sheet, auctions, latest_purchases, source: str) -> None:
    headers = [
        "Номер тендера",
        "Название",
        "Заказчик",
        "НМЦ",
        "fit_status",
        "Выгодность",
        "Маржа руб",
        "Маржа %",
        "Рейтинг",
        "Период подачи",
        "Дедлайн",
        "Ссылка на карточку",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for auction in auctions:
        purchase = latest_purchases.get(auction.auction_number)
        date_range = purchase.date_range if purchase else ""
        deadline = parse_deadline_from_date_range(date_range)
        deadline_text = deadline.strftime("%d.%m.%Y %H:%M") if deadline else ""

        if source == "yandex":
            margin_rub = auction.margin_rub_yandex
            margin_pct = auction.margin_percent_yandex
            rating = auction.rating_yandex
            fit_status = auction.fit_status_yandex
        else:
            margin_rub = auction.margin_rub
            margin_pct = auction.margin_percent
            rating = auction.rating
            fit_status = auction.fit_status

        sheet.append(
            [
                auction.auction_number,
                auction.name,
                auction.customer,
                float_or_none(auction.initial_price),
                fit_status,
                fit_status_label(fit_status),
                float_or_none(margin_rub),
                float_or_none(margin_pct),
                rating,
                date_range,
                deadline_text,
                auction.detail_url,
            ]
        )

    autosize_columns(sheet)


def write_positions_sheet(sheet, auctions, items_by_auction, latest_purchases, source: str) -> None:
    if source == "yandex":
        headers = [
            "Номер тендера",
            "Название тендера",
            "fit_status",
            "Выгодность",
            "Маржа руб (yandex)",
            "Маржа %",
            "Ссылка на тендер",
            "№ позиции",
            "Наименование",
            "Кол-во",
            "Ед.",
            "Цена в тендере за ед.",
            "Сумма в тендере",
            "Статус Yandex",
            "Мин",
            "Макс",
            "Медиана",
            "Сумма закупки (медиана)",
            "Тендер vs рынок",
            "Комментарий",
            "Карточки (JSON)",
        ]
    else:
        headers = [
            "Номер тендера",
            "Название тендера",
            "fit_status",
            "Выгодность",
            "Маржа руб",
            "Маржа %",
            "Ссылка на тендер",
            "№ позиции",
            "Наименование",
            "Кол-во",
            "Ед.",
            "Цена в тендере за ед.",
            "Сумма в тендере",
            "Ссылка SKU",
            "Статус цены",
            "Рынок за ед.",
            "Сумма закупки (оценка)",
            "Уверенность",
            "Тип",
            "Комментарий",
            "Ссылки на источники",
            "Ход агента",
        ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for auction in auctions:
        items = items_by_auction.get(auction.auction_number, [])

        if source == "yandex":
            margin_rub = float_or_none(auction.margin_rub_yandex)
            margin_pct = float_or_none(auction.margin_percent_yandex)
            fit_status = auction.fit_status_yandex
        else:
            margin_rub = float_or_none(auction.margin_rub)
            margin_pct = float_or_none(auction.margin_percent)
            fit_status = auction.fit_status

        fit_text = fit_status_label(fit_status)

        for item in items:
            quantity = float_or_none(item.quantity)
            price_per_unit = float_or_none(item.price_per_unit)
            total_price = float_or_none(item.total_price)
            tender_line = tender_line_total(total_price, price_per_unit, quantity)

            if source == "yandex":
                median = float_or_none(item.yandex_price_median)
                eval_status = item.yandex_research_status
                if eval_status == "done" and median is not None:
                    eval_status = "priced"
                elif eval_status != "pending":
                    eval_status = "not_found"
                market_line_cost = line_cost_for_item(
                    eval_status if eval_status != "done" else "priced",
                    tender_line,
                    quantity,
                    median,
                )
                sheet.append(
                    [
                        auction.auction_number,
                        auction.name,
                        fit_status,
                        fit_text,
                        margin_rub,
                        margin_pct,
                        auction.detail_url,
                        item.position_index,
                        item.name,
                        quantity,
                        item.unit,
                        price_per_unit,
                        tender_line,
                        item.yandex_research_status,
                        float_or_none(item.yandex_price_min),
                        float_or_none(item.yandex_price_max),
                        median,
                        market_line_cost,
                        item.yandex_tender_vs_market,
                        item.yandex_note,
                        item.yandex_cards_json,
                    ]
                )
                continue

            market_per_unit = float_or_none(item.market_price_per_unit)
            market_line_cost = line_cost_for_item(
                item.price_research_status,
                tender_line,
                quantity,
                market_per_unit,
            )

            sheet.append(
                [
                    auction.auction_number,
                    auction.name,
                    fit_status,
                    fit_text,
                    margin_rub,
                    margin_pct,
                    auction.detail_url,
                    item.position_index,
                    item.name,
                    quantity,
                    item.unit,
                    price_per_unit,
                    tender_line,
                    item.sku_url,
                    item.price_research_status,
                    market_per_unit,
                    market_line_cost,
                    item.market_confidence,
                    item.market_item_type,
                    item.market_note,
                    sources_to_urls(item.market_sources_json),
                    excel_agent_log(item.market_agent_log),
                ]
            )
            log_cell = sheet.cell(sheet.max_row, len(headers))
            log_cell.alignment = Alignment(wrap_text=True, vertical="top")

    autosize_columns(sheet)


def _rows_below_header(sheet) -> list[list]:
    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        rows.append(list(row))
    return rows


def _drop_auction_rows(sheet, auction_numbers: set[str]) -> None:
    for row_index in range(sheet.max_row, 1, -1):
        cell_value = sheet.cell(row_index, 1).value
        if cell_value is not None and str(cell_value) in auction_numbers:
            sheet.delete_rows(row_index)


def _append_rows(sheet, rows: list[list]) -> None:
    for row in rows:
        sheet.append(row)


def upsert_claude_auctions(output_path: Path, auction_numbers: list[str]) -> int:
    """
    Дописывает или обновляет строки этих тендеров в уже существующем Excel.
    Остальные номера в файле не трогает.
    """
    auctions, items_by_auction, latest_purchases = load_tenders_for_export(
        "claude",
        auction_numbers,
    )
    if not auctions:
        return 0

    fresh = Workbook()
    summary_sheet = fresh.active
    summary_sheet.title = "Тендеры"
    write_summary_sheet(summary_sheet, auctions, latest_purchases, "claude")
    positions_sheet = fresh.create_sheet("Позиции")
    write_positions_sheet(positions_sheet, auctions, items_by_auction, latest_purchases, "claude")

    summary_rows = _rows_below_header(summary_sheet)
    position_rows = _rows_below_header(positions_sheet)
    numbers = {auction.auction_number for auction in auctions}

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = None
    if output_path.is_file() and output_path.stat().st_size > 0:
        try:
            workbook = load_workbook(output_path)
        except BadZipFile:
            workbook = None
    if workbook is None:
        workbook = fresh
        autosize_columns(workbook["Тендеры"])
        autosize_columns(workbook["Позиции"])
        workbook.save(output_path)
        return len(auctions)

    if "Тендеры" not in workbook.sheetnames or "Позиции" not in workbook.sheetnames:
        raise SystemExit(f"В {output_path.name} нет листов «Тендеры» и «Позиции»")

    positions = workbook["Позиции"]
    if positions.cell(1, positions.max_column).value != "Ход агента":
        header_cell = positions.cell(1, positions.max_column + 1, "Ход агента")
        header_cell.font = Font(bold=True)

    _drop_auction_rows(workbook["Тендеры"], numbers)
    _append_rows(workbook["Тендеры"], summary_rows)
    _drop_auction_rows(workbook["Позиции"], numbers)
    _append_rows(workbook["Позиции"], position_rows)
    autosize_columns(workbook["Тендеры"])
    autosize_columns(workbook["Позиции"])
    workbook.save(output_path)
    return len(auctions)


def export_to_excel(output_path: Path, source: str) -> int:
    auctions, items_by_auction, latest_purchases = load_tenders_for_export(source)
    if not auctions:
        return 0

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Тендеры"
    write_summary_sheet(summary_sheet, auctions, latest_purchases, source)

    positions_sheet = workbook.create_sheet("Позиции")
    write_positions_sheet(positions_sheet, auctions, items_by_auction, latest_purchases, source)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return len(auctions)


def main() -> None:
    parser = argparse.ArgumentParser(description="Excel: тендеры с оценкой цен и все позиции")
    parser.add_argument(
        "--source",
        choices=("claude", "yandex"),
        default="claude",
        help="claude = market_*, yandex = медиана вкладки Товары",
    )
    parser.add_argument("--out", default=None, help="Путь к .xlsx")
    args = parser.parse_args()

    if args.out:
        output_path = Path(args.out)
    elif args.source == "yandex":
        output_path = DEFAULT_OUTPUT_YANDEX
    else:
        output_path = DEFAULT_OUTPUT_CLAUDE

    count = export_to_excel(output_path, args.source)

    if count == 0:
        print("Нет тендеров с оценёнными ценами (хотя бы одна позиция не pending).")
        return

    print(f"Сохранено тендеров: {count}")
    print(f"Файл: {output_path.resolve()}")
    print(f"Обновлено: {datetime.now(timezone.utc).astimezone().strftime('%Y-%m-%d %H:%M %Z')}")


if __name__ == "__main__":
    main()
